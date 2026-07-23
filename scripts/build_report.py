from __future__ import annotations

import argparse
from datetime import date
import json
from pathlib import Path
from typing import Any, Iterable

from docx import Document
from docx.enum.table import WD_CELL_VERTICAL_ALIGNMENT, WD_TABLE_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Cm, Inches, Pt, RGBColor

PROJECT_ROOT = Path(__file__).resolve().parents[1]
DEFAULT_TEMPLATE = PROJECT_ROOT / "docs/templates/rapport-technique-template.docx"
DEFAULT_OUTPUT = PROJECT_ROOT / "outputs/rapport-technique-openagenda-rag.docx"

NAVY = "18324A"
TEAL = "147D80"
ORANGE = "E77C3C"
CREAM = "F7F2E8"
LIGHT_TEAL = "E7F2F1"
LIGHT_ORANGE = "FBEADF"
WHITE = "FFFFFF"
GRAY = RGBColor(74, 85, 95)


def parse_args() -> argparse.Namespace:
    parser = argparse.ArgumentParser(
        description="Build the OpenAgenda RAG technical report from the supplied template."
    )
    parser.add_argument("--template", type=Path, default=DEFAULT_TEMPLATE)
    parser.add_argument("--output", type=Path, default=DEFAULT_OUTPUT)
    parser.add_argument(
        "--ragas-results",
        type=Path,
        default=PROJECT_ROOT / "outputs/evaluation/ragas_results.json",
    )
    parser.add_argument(
        "--fetch-manifest",
        type=Path,
        default=PROJECT_ROOT / "seed-data/processed/fetch_manifest.json",
    )
    parser.add_argument(
        "--index-manifest",
        type=Path,
        default=PROJECT_ROOT / "seed-data/index/index_manifest.json",
    )
    return parser.parse_args()


def _clear_body(document: Document) -> None:
    body = document._element.body
    for child in list(body):
        if child.tag != qn("w:sectPr"):
            body.remove(child)


def _set_cell_shading(cell: Any, fill: str) -> None:
    properties = cell._tc.get_or_add_tcPr()
    shading = properties.find(qn("w:shd"))
    if shading is None:
        shading = OxmlElement("w:shd")
        properties.append(shading)
    shading.set(qn("w:fill"), fill)


def _set_cell_margins(cell: Any, value: int = 120) -> None:
    properties = cell._tc.get_or_add_tcPr()
    margins = properties.first_child_found_in("w:tcMar")
    if margins is None:
        margins = OxmlElement("w:tcMar")
        properties.append(margins)
    for edge in ("top", "left", "bottom", "right"):
        node = margins.find(qn(f"w:{edge}"))
        if node is None:
            node = OxmlElement(f"w:{edge}")
            margins.append(node)
        node.set(qn("w:w"), str(value))
        node.set(qn("w:type"), "dxa")


def _set_table_borders(table: Any, color: str = "D5DEE3") -> None:
    properties = table._tbl.tblPr
    borders = properties.first_child_found_in("w:tblBorders")
    if borders is None:
        borders = OxmlElement("w:tblBorders")
        properties.append(borders)
    for edge in ("top", "left", "bottom", "right", "insideH", "insideV"):
        node = borders.find(qn(f"w:{edge}"))
        if node is None:
            node = OxmlElement(f"w:{edge}")
            borders.append(node)
        node.set(qn("w:val"), "single")
        node.set(qn("w:sz"), "6")
        node.set(qn("w:color"), color)


def _add_page_number(paragraph: Any) -> None:
    run = paragraph.add_run()
    begin = OxmlElement("w:fldChar")
    begin.set(qn("w:fldCharType"), "begin")
    instruction = OxmlElement("w:instrText")
    instruction.set(qn("xml:space"), "preserve")
    instruction.text = " PAGE "
    end = OxmlElement("w:fldChar")
    end.set(qn("w:fldCharType"), "end")
    run._r.extend([begin, instruction, end])


def _configure_document(document: Document) -> None:
    section = document.sections[0]
    section.top_margin = Cm(1.7)
    section.bottom_margin = Cm(1.5)
    section.left_margin = Cm(1.8)
    section.right_margin = Cm(1.8)

    styles = document.styles
    normal = styles["normal"]
    normal.font.name = "Aptos"
    normal.font.size = Pt(10)
    normal.font.color.rgb = GRAY
    normal.paragraph_format.space_after = Pt(5)
    normal.paragraph_format.line_spacing = 1.08

    for style_name, size, color in (
        ("Title", 30, NAVY),
        ("Heading 1", 25, NAVY),
        ("Heading 2", 17, TEAL),
        ("Heading 3", 12, ORANGE),
    ):
        style = styles[style_name]
        style.font.name = "Aptos Display"
        style.font.size = Pt(size)
        style.font.color.rgb = RGBColor.from_string(color)
        style.font.bold = True
        style.paragraph_format.space_before = Pt(8)
        style.paragraph_format.space_after = Pt(6)

    footer = section.footer.paragraphs[0]
    footer.alignment = WD_ALIGN_PARAGRAPH.CENTER
    footer.add_run("Puls-Events | Assistant RAG OpenAgenda | ")
    _add_page_number(footer)


def _add_title_page(document: Document, fetch: dict[str, Any], index: dict[str, Any]) -> None:
    spacer = document.add_paragraph()
    spacer.paragraph_format.space_after = Pt(55)

    title = document.add_paragraph(style="Title")
    title.alignment = WD_ALIGN_PARAGRAPH.LEFT
    title.add_run("Assistant intelligent de recommandation\nd'événements culturels")

    subtitle = document.add_paragraph()
    subtitle.paragraph_format.space_before = Pt(10)
    run = subtitle.add_run("Rapport technique du POC RAG pour Puls-Events")
    run.font.name = "Aptos Display"
    run.font.size = Pt(17)
    run.font.bold = True
    run.font.color.rgb = RGBColor.from_string(TEAL)

    table = document.add_table(rows=1, cols=3)
    table.alignment = WD_TABLE_ALIGNMENT.CENTER
    table.autofit = False
    labels = [
        ("CORPUS", f"{fetch['normalized_event_count']:,}".replace(",", " ") + " événements"),
        ("INDEX", f"{index['indexed_document_count']:,}".replace(",", " ") + " chunks"),
        ("ZONE", fetch["filters"]["city"]),
    ]
    for cell, (label, value) in zip(table.rows[0].cells, labels, strict=True):
        cell.width = Inches(2.0)
        cell.vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER
        _set_cell_shading(cell, LIGHT_TEAL)
        _set_cell_margins(cell, 180)
        paragraph = cell.paragraphs[0]
        paragraph.alignment = WD_ALIGN_PARAGRAPH.CENTER
        label_run = paragraph.add_run(label + "\n")
        label_run.bold = True
        label_run.font.size = Pt(9)
        label_run.font.color.rgb = RGBColor.from_string(TEAL)
        value_run = paragraph.add_run(value)
        value_run.bold = True
        value_run.font.size = Pt(14)
        value_run.font.color.rgb = RGBColor.from_string(NAVY)

    document.add_paragraph()
    details = document.add_paragraph()
    details.paragraph_format.space_before = Pt(25)
    details.add_run("Technologies : ").bold = True
    details.add_run("Python 3.12, LangChain, Mistral, FAISS, FastAPI, Docker, RAGAS")
    details.add_run("\nCorpus : ").bold = True
    filters = fetch["filters"]
    details.add_run(
        f"Paris, du {filters['start_date']} au {filters['end_date']}, source OpenAgenda/Opendatasoft"
    )
    details.add_run("\nDate du rapport : ").bold = True
    details.add_run(date.today().strftime("%d/%m/%Y"))

    note = document.add_paragraph()
    note.paragraph_format.space_before = Pt(45)
    note.add_run(
        "Ce document complète le modèle de rapport fourni et décrit uniquement "
        "les composants et résultats vérifiés dans le dépôt."
    ).italic = True
    document.add_page_break()


def _add_heading(document: Document, number: int, title: str) -> None:
    document.add_heading(f"{number}. {title}", level=1)


def _add_bullets(document: Document, items: Iterable[str]) -> None:
    for item in items:
        paragraph = document.add_paragraph()
        paragraph.paragraph_format.left_indent = Cm(0.45)
        paragraph.paragraph_format.first_line_indent = Cm(-0.25)
        bullet = paragraph.add_run("• ")
        bullet.bold = True
        bullet.font.color.rgb = RGBColor.from_string(ORANGE)
        paragraph.add_run(item)


def _add_code(document: Document, text: str) -> None:
    paragraph = document.add_paragraph()
    paragraph.paragraph_format.left_indent = Cm(0.4)
    paragraph.paragraph_format.right_indent = Cm(0.4)
    paragraph.paragraph_format.space_before = Pt(4)
    paragraph.paragraph_format.space_after = Pt(7)
    run = paragraph.add_run(text)
    run.font.name = "Liberation Mono"
    run.font.size = Pt(8.5)
    run.font.color.rgb = RGBColor.from_string(NAVY)
    properties = paragraph._p.get_or_add_pPr()
    shading = OxmlElement("w:shd")
    shading.set(qn("w:fill"), CREAM)
    properties.append(shading)


def _add_pipeline_table(document: Document) -> None:
    stages = [
        ("1", "Collecte", "Opendatasoft v2.1"),
        ("2", "Préparation", "Pandas + Parquet"),
        ("3", "Indexation", "Mistral + FAISS"),
        ("4", "RAG", "LangChain + Mistral"),
        ("5", "Exposition", "FastAPI + Docker"),
    ]
    table = document.add_table(rows=1, cols=len(stages))
    table.alignment = WD_TABLE_ALIGNMENT.CENTER
    table.autofit = True
    for cell, (number, title, detail) in zip(table.rows[0].cells, stages, strict=True):
        _set_cell_shading(cell, LIGHT_TEAL if int(number) % 2 else LIGHT_ORANGE)
        _set_cell_margins(cell, 100)
        paragraph = cell.paragraphs[0]
        paragraph.alignment = WD_ALIGN_PARAGRAPH.CENTER
        first = paragraph.add_run(number + "\n")
        first.bold = True
        first.font.size = Pt(15)
        first.font.color.rgb = RGBColor.from_string(ORANGE)
        second = paragraph.add_run(title + "\n")
        second.bold = True
        second.font.size = Pt(9)
        second.font.color.rgb = RGBColor.from_string(NAVY)
        detail_run = paragraph.add_run(detail)
        detail_run.font.size = Pt(7.5)


def _add_metric_table(document: Document, ragas: dict[str, Any]) -> None:
    labels = {
        "faithfulness": "Fidélité au contexte",
        "answer_relevancy": "Pertinence de la réponse",
        "context_precision": "Précision du contexte",
        "answer_correctness": "Correction de la réponse",
    }
    summary = ragas["summary"]
    table = document.add_table(rows=1, cols=3)
    table.alignment = WD_TABLE_ALIGNMENT.CENTER
    _set_table_borders(table)
    for cell, value in zip(
        table.rows[0].cells,
        ("Métrique", "Score moyen", "Interprétation"),
        strict=True,
    ):
        _set_cell_shading(cell, NAVY)
        run = cell.paragraphs[0].add_run(value)
        run.bold = True
        run.font.color.rgb = RGBColor.from_string(WHITE)
    for name in ragas["metric_names"]:
        cells = table.add_row().cells
        score = summary["metric_means"][name]
        values = (labels[name], f"{score:.3f}", summary["metric_interpretations"][name])
        for cell, value in zip(cells, values, strict=True):
            cell.text = value
            _set_cell_margins(cell)
        _set_cell_shading(cells[1], LIGHT_TEAL if score >= 0.8 else LIGHT_ORANGE)


def _add_ragas_examples(document: Document, ragas: dict[str, Any]) -> None:
    for index, example in enumerate(ragas["examples"], start=1):
        if index > 1:
            document.add_page_break()
        document.add_heading(f"Exemple RAGAS {index} — {example['case_id']}", level=2)
        document.add_paragraph(f"Question : {example['question']}")
        context = example["retrieved_contexts"][0].replace("\n", " ")
        document.add_paragraph(
            "Contexte principal : " + context[:650] + ("…" if len(context) > 650 else "")
        )
        document.add_paragraph("Réponse générée : " + example["response"])
        table = document.add_table(rows=1, cols=4)
        _set_table_borders(table)
        for cell, name in zip(table.rows[0].cells, ragas["metric_names"], strict=True):
            _set_cell_shading(cell, NAVY)
            run = cell.paragraphs[0].add_run(name.replace("_", " "))
            run.bold = True
            run.font.size = Pt(8)
            run.font.color.rgb = RGBColor.from_string(WHITE)
        score_row = table.add_row().cells
        for cell, name in zip(score_row, ragas["metric_names"], strict=True):
            score = example["scores"][name]
            cell.text = f"{score:.3f}\n{example['interpretations'][name]}"
            _set_cell_shading(cell, LIGHT_TEAL if score >= 0.8 else LIGHT_ORANGE)
        document.add_paragraph(
            f"Interprétation : score global {example['overall_score']:.3f}, "
            f"niveau « {example['overall_interpretation']} ». "
            "Les écarts orientent les améliorations du prompt et des références.",
        )


def build_report(
    template_path: Path,
    output_path: Path,
    ragas_path: Path,
    fetch_manifest_path: Path,
    index_manifest_path: Path,
) -> Path:
    for path in (template_path, ragas_path, fetch_manifest_path, index_manifest_path):
        if not path.is_file():
            raise FileNotFoundError(path)

    ragas = json.loads(ragas_path.read_text(encoding="utf-8"))
    fetch = json.loads(fetch_manifest_path.read_text(encoding="utf-8"))
    index = json.loads(index_manifest_path.read_text(encoding="utf-8"))

    document = Document(template_path)
    _clear_body(document)
    _configure_document(document)
    _add_title_page(document, fetch, index)

    _add_heading(document, 1, "Objectifs du projet")
    document.add_paragraph(
        "Puls-Events souhaite tester un assistant culturel capable de répondre à "
        "des demandes naturelles en s'appuyant sur des événements réels. Le POC "
        "démontre la faisabilité technique, la traçabilité des recommandations et "
        "l'exposition du service aux équipes produit et marketing."
    )
    _add_bullets(
        document,
        [
            "Besoin : recommandations pertinentes, formulées naturellement et vérifiables.",
            "Réponse : recherche sémantique dans FAISS, puis génération augmentée par Mistral.",
            "Périmètre livré : Paris, du 10 juillet 2025 au 10 juillet 2026.",
            f"Volume : {fetch['normalized_event_count']} événements normalisés et {index['indexed_document_count']} chunks.",
        ],
    )
    document.add_page_break()

    _add_heading(document, 2, "Architecture du système")
    _add_pipeline_table(document)
    document.add_heading("Séparation des responsabilités", level=2)
    document.add_paragraph(
        "Les scripts publics orchestrent des modules dédiés à l'ingestion, "
        "l'indexation, la génération RAG, l'API et l'évaluation. Le service "
        "`OpenAgendaRAGService` conserve le retriever et le modèle en mémoire."
    )
    _add_bullets(
        document,
        [
            "Entrée : API Opendatasoft v2.1, dataset evenements-publics-openagenda.",
            "Stockage intermédiaire : JSON brut, Parquet normalisé et manifestes.",
            "Recherche : embeddings mistral-embed et index local FAISS persistant.",
            "Génération : mistral-small-latest orchestré par LangChain.",
            "Livraison : API FastAPI, documentation Swagger et conteneur Docker.",
        ],
    )
    document.add_page_break()

    _add_heading(document, 3, "Préparation et vectorisation des données")
    document.add_heading("Collecte et filtres", level=2)
    _add_bullets(
        document,
        [
            f"Ville : {fetch['filters']['city']}.",
            f"Fenêtre : {fetch['filters']['start_date']} à {fetch['filters']['end_date']}.",
            "Pagination complète, déduplication par identifiant et gestion des champs absents.",
            "Schéma stable incluant texte, ville, lieu, coordonnées, dates, URL et catégories.",
        ],
    )
    document.add_heading("Chunking et embeddings", level=2)
    document.add_paragraph(
        f"Chaque `text_for_embedding` est découpé en chunks de {index['chunk_size']} "
        f"caractères avec {index['chunk_overlap']} caractères de chevauchement. "
        f"Les vecteurs sont produits par `{index['embedding_model']}` par lots de "
        f"{index['batch_size']}."
    )
    _add_code(
        document,
        "python scripts/fetch_events.py --city Paris --start-date 2025-07-10 "
        "--end-date 2026-07-10\npython scripts/build_index.py --rebuild",
    )
    document.add_page_break()

    _add_heading(document, 4, "Choix du modèle NLP")
    _add_bullets(
        document,
        [
            "Embeddings : mistral-embed, adapté à la similarité sémantique et intégré à LangChain.",
            "Génération : mistral-small-latest, compromis coût, latence et qualité pour le POC.",
            "Prompt : réponse en français, fondée uniquement sur le contexte, avec refus explicite si l'information manque.",
            "Température basse et graine fixe pour stabiliser les démonstrations.",
        ],
    )
    document.add_heading("Limites", level=2)
    document.add_paragraph(
        "La génération et l'évaluation dépendent de l'API Mistral. Le compte gratuit "
        "impose une exécution RAGAS séquentielle avec reprise automatique après une "
        "erreur temporaire. Les sorties restent non déterministes au mot près."
    )

    _add_heading(document, 5, "Construction de la base vectorielle")
    document.add_paragraph(
        "FAISS est retenu pour sa rapidité, sa portabilité et sa persistance locale. "
        "L'index officiel est sauvegardé dans `data/index/faiss/` et accompagné "
        "d'un manifeste et d'un parquet d'inspection."
    )
    _add_bullets(
        document,
        [
            f"{index['indexed_event_count']} événements indexés.",
            f"{index['indexed_document_count']} documents vectoriels.",
            "Métadonnées : identifiants, titre, ville, lieu, dates, fuseau, URL, catégories et date source.",
            "Recherche top-k par similarité ; rechargement local testé après sauvegarde.",
        ],
    )
    document.add_page_break()

    _add_heading(document, 6, "API et endpoints exposés")
    _add_bullets(
        document,
        [
            "`GET /health` : état du service et paramètres principaux.",
            "`POST /ask` : question non vide, réponse augmentée et sources.",
            "`POST /rebuild` : reconstruction complète puis réinitialisation du runtime.",
            "`GET /docs` : documentation Swagger générée par FastAPI.",
        ],
    )
    _add_code(
        document,
        'curl -X POST http://localhost:8000/ask \\\n'
        '  -H "Content-Type: application/json" \\\n'
        '  -d \'{"question":"Parle-moi de Concert Fishers a Paris"}\'',
    )
    document.add_paragraph(
        "Les questions vides et les erreurs de configuration produisent des réponses "
        "HTTP explicites. La clé Mistral n'est jamais retournée. `/rebuild` peut être "
        "protégé par l'en-tête `X-Admin-Token`."
    )
    document.add_page_break()

    _add_heading(document, 7, "Évaluation du système")
    document.add_paragraph(
        f"Le jeu versionné contient {ragas['summary']['example_count']} cas annotés. "
        "Chaque exemple conserve la question, les contextes exactement transmis au "
        "LLM, la réponse, une référence humaine, quatre scores et leur interprétation."
    )
    _add_metric_table(document, ragas)
    document.add_paragraph(
        "Seuils d'interprétation : fort ≥ 0,80 ; acceptable ≥ 0,60 ; à améliorer < 0,60. "
        "Trois exemples démontrent la méthode mais ne constituent pas un benchmark statistique."
    )
    document.add_page_break()
    _add_ragas_examples(document, ragas)
    document.add_page_break()

    _add_heading(document, 8, "Recommandations et perspectives")
    _add_bullets(
        document,
        [
            "Élargir le jeu annoté à plusieurs villes, catégories, saisons et cas négatifs.",
            "Ajouter un seuil de similarité afin de rejeter plus tôt les questions hors corpus.",
            "Renforcer le prompt pour mieux aligner la réponse avec la référence attendue.",
            "Automatiser la collecte, la reconstruction et l'évaluation dans une CI.",
            "Ajouter authentification, supervision, limites de débit et tests de charge avant production.",
        ],
    )
    document.add_paragraph(
        "Le POC valide la faisabilité : le retrieval obtient une précision de contexte "
        f"de {ragas['summary']['metric_means']['context_precision']:.3f} et la démonstration "
        "API complète s'exécute en moins de cinq minutes."
    )

    _add_heading(document, 9, "Organisation du dépôt GitHub")
    _add_code(
        document,
        "src/openagenda_rag/    logique ingestion, indexation, RAG, API, évaluation\n"
        "scripts/               commandes publiques et démonstration\n"
        "tests/                 tests réseau-off et fixtures annotées\n"
        "data/                  artefacts bruts, structurés et index FAISS\n"
        "outputs/               résultats RAGAS, chronométrage, PPTX et rapport\n"
        "docs/soutenance/       script oral et questions-réponses\n"
        "Dockerfile             image de l'API\n"
        "docker-compose.yml     exécution locale avec volume persistant",
    )
    document.add_paragraph(
        "`environment.yml` est la source Conda et `requirements.txt` sert aux "
        "installations pip et Docker. `.env` et les secrets sont ignorés par Git."
    )
    document.add_page_break()

    _add_heading(document, 10, "Annexes")
    document.add_heading("Commandes de reproduction", level=2)
    _add_code(
        document,
        "conda env create -f environment.yml\n"
        "conda activate openagenda-rag\n"
        "python scripts/check_env.py\n"
        "python scripts/evaluate_ragas.py\n"
        "docker compose up --build\n"
        "bash scripts/demo_api_5min.sh",
    )
    document.add_heading("Traçabilité de l'évaluation", level=2)
    provenance = ragas.get("provenance", {})
    _add_bullets(
        document,
        [
            f"RAGAS : version {provenance.get('ragas_version', 'non renseignée')}.",
            f"Jeu annoté SHA-256 : {provenance.get('dataset_sha256', 'non renseigné')}.",
            f"Manifeste FAISS SHA-256 : {provenance.get('index_manifest_sha256', 'non renseigné')}.",
            f"Commit de l'évaluation : {provenance.get('git_commit', 'non renseigné')}.",
            f"Modèle juge : {ragas['models']['chat']}.",
            f"Modèle d'embedding : {ragas['models']['embedding']}.",
            f"Évaluation générée le {ragas['generated_at']}.",
        ],
    )
    document.add_heading("Livrables de soutenance", level=2)
    _add_bullets(
        document,
        [
            "`outputs/openagenda-rag-soutenance.pptx` : présentation de 15 slides.",
            "`docs/soutenance/script_soutenance_15min.md` : conducteur oral de 15 minutes.",
            "`docs/soutenance/questions_reponses.md` : 14 réponses techniques et métier.",
            "`outputs/demo/demo_api_timing.json` : preuve de démonstration chronométrée.",
        ],
    )

    output_path.parent.mkdir(parents=True, exist_ok=True)
    document.save(output_path)
    return output_path


def main() -> int:
    args = parse_args()
    output = build_report(
        template_path=args.template,
        output_path=args.output,
        ragas_path=args.ragas_results,
        fetch_manifest_path=args.fetch_manifest,
        index_manifest_path=args.index_manifest,
    )
    print(output)
    return 0


if __name__ == "__main__":
    raise SystemExit(main())
